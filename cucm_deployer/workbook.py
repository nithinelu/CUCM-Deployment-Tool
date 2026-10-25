"""Excel workbook: build the template, read rows, write results back.

Every object sheet has the same layout:
    row 1  headers   (Action | fields... | Status | Message | Updated)
    row 2  hints     (grey, ignored when reading)
    row 3+ data      (one CUCM object per row)
"""

import datetime as dt
import os

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter, quote_sheetname
from openpyxl.worksheet.datavalidation import DataValidation

from .objects import ACTIONS, BOOL_CHOICES, DEPLOY_PLAN, SHEET_BY_NAME, SHEETS

DONE_STATUSES = ("Success", "Updated", "Deleted", "Exists")
from .schema import get_schema

HEADER_ROW, HINT_ROW, DATA_ROW = 1, 2, 3
STATUS_COLS = ["Status", "Message", "Updated"]
MAX_ROWS = 5000
EXTRA_PREFIX = "axl:"

NAVY = "1F3864"
TEAL = "0E7C86"
RED = "9C2A2A"
GREY = "7F7F7F"
thin = Side(style="thin", color="D0D7E2")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)


class Row:
    def __init__(self, sheet, excel_row, action, cells, extra, status):
        self.sheet = sheet          # sheet name
        self.excel_row = excel_row
        self.action = action        # Add | Update | Add/Update | Delete | Skip
        self.cells = cells          # header -> raw value (only non-blank)
        self.extra = extra          # "axl:" path -> raw value
        self.status = status        # Status column from a previous run

    def label(self):
        sh = SHEET_BY_NAME[self.sheet]
        keys = [str(self.cells.get(f.header, "")) for f in sh.fields if f.key]
        return " / ".join(k for k in keys if k) or f"row {self.excel_row}"


def norm(header):
    return str(header or "").strip().rstrip("*").strip().lower()


# --------------------------------------------------------------------------
# reading
# --------------------------------------------------------------------------
def read_workbook(path):
    """-> {sheet_name: [Row, ...]} for every known sheet present in the file."""
    wb = load_workbook(path, data_only=True)
    out = {}
    for sh in SHEETS:
        if sh.name not in wb.sheetnames:
            continue
        ws = wb[sh.name]
        headers = [c.value for c in ws[HEADER_ROW]]
        cols = {}
        for idx, h in enumerate(headers):
            n = norm(h)
            if not n:
                continue
            if n.startswith(EXTRA_PREFIX):
                cols[idx] = ("extra", str(h).strip()[len(EXTRA_PREFIX):].strip())
            elif n == "action":
                cols[idx] = ("action", None)
            elif n == "status":
                cols[idx] = ("status", None)
            else:
                f = sh.field(str(h))
                if f is not None:
                    cols[idx] = ("field", f.header)
        rows = []
        for r_idx, values in enumerate(ws.iter_rows(min_row=DATA_ROW, values_only=True),
                                       start=DATA_ROW):
            cells, extra, action, status = {}, {}, "", ""
            for idx, v in enumerate(values):
                if idx not in cols or v is None or (isinstance(v, str) and not v.strip()):
                    continue
                kind, name = cols[idx]
                if kind == "field":
                    cells[name] = v
                elif kind == "extra":
                    extra[name] = v
                elif kind == "action":
                    action = str(v).strip()
                elif kind == "status":
                    status = str(v).strip()
            if not cells and not extra:
                continue
            rows.append(Row(sh.name, r_idx, normalise_action(action), cells, extra, status))
        out[sh.name] = rows
    return out


def normalise_action(text):
    t = (text or "").strip().lower().replace(" ", "")
    for a in ACTIONS:
        if t == a.lower().replace(" ", ""):
            return a
    if t in ("", "new", "create", "a"):
        return "Add"
    if t in ("upsert", "add/upd", "addorupdate", "merge"):
        return "Add/Update"
    if t in ("remove", "del", "d"):
        return "Delete"
    if t in ("modify", "change", "u"):
        return "Update"
    return text or "Add"        # unknown -> reported by validation


def read_settings(path):
    """Cluster settings from the Dashboard sheet (no passwords are stored)."""
    try:
        wb = load_workbook(path, data_only=True, read_only=True)
    except Exception:
        return {}
    if "Dashboard" not in wb.sheetnames:
        return {}
    ws = wb["Dashboard"]
    out = {}
    for row in ws.iter_rows(min_row=1, max_row=12, min_col=2, max_col=3, values_only=True):
        if row and row[0] in ("CUCM Publisher", "AXL User", "AXL Version"):
            out[row[0]] = "" if row[1] is None else str(row[1]).strip()
    wb.close()
    return out


# --------------------------------------------------------------------------
# writing results
# --------------------------------------------------------------------------
def write_results(path, results, out_path=None):
    """results: iterable of (sheet, excel_row, status, message).
    Saves to out_path (default: same file). Returns the path written."""
    wb = load_workbook(path)
    now = dt.datetime.now().strftime("%Y-%m-%d %H:%M")
    col_cache = {}
    for sheet, excel_row, status, message in results:
        if sheet not in wb.sheetnames:
            continue
        ws = wb[sheet]
        if sheet not in col_cache:
            headers = {norm(c.value): c.column for c in ws[HEADER_ROW] if c.value}
            nxt = ws.max_column + 1
            cols = []
            for name in STATUS_COLS:
                if name.lower() not in headers:
                    ws.cell(HEADER_ROW, nxt, name)
                    headers[name.lower()] = nxt
                    nxt += 1
                cols.append(headers[name.lower()])
            col_cache[sheet] = cols
        c_status, c_msg, c_time = col_cache[sheet]
        if status == "Exists" and ws.cell(excel_row, c_status).value == "Success":
            # created by an earlier run - keep it ours so a rollback still deletes it
            status, message = "Success", "Already deployed by an earlier run"
        ws.cell(excel_row, c_status, status)
        ws.cell(excel_row, c_msg, (message or "")[:1000])
        ws.cell(excel_row, c_time, now)
    target = out_path or path
    wb.save(target)
    return target


def make_rollback(path, out_path):
    """Copy of the workbook where every row that was successfully *added*
    becomes a Delete row (run it to undo a deployment)."""
    wb = load_workbook(path)
    count = 0
    for sh in SHEETS:
        if sh.name not in wb.sheetnames:
            continue
        ws = wb[sh.name]
        headers = {norm(c.value): c.column for c in ws[HEADER_ROW] if c.value}
        a_col, s_col = headers.get("action"), headers.get("status")
        if not a_col:
            continue
        for r in range(DATA_ROW, ws.max_row + 1):
            if sh.kind == "relations" or not s_col:
                ws.cell(r, a_col, "Skip")
                continue
            status = str(ws.cell(r, s_col).value or "")
            action = normalise_action(str(ws.cell(r, a_col).value or ""))
            if status == "Success" and action in ("Add", "Add/Update"):
                ws.cell(r, a_col, "Delete")
                ws.cell(r, s_col).value = None
                count += 1
            else:
                ws.cell(r, a_col, "Skip")
    wb.save(out_path)
    return count


# --------------------------------------------------------------------------
# template
# --------------------------------------------------------------------------
def _hdr(cell, fill, color="FFFFFF"):
    cell.font = Font(bold=True, color=color)
    cell.fill = PatternFill("solid", fgColor=fill)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = BORDER


def build_template(path, sample=None):
    """Write a new workbook. sample: {sheet: [{header: value}]} example rows."""
    schema = get_schema()
    wb = Workbook()
    dash = wb.active
    dash.title = "Dashboard"
    lists = wb.create_sheet("Lists")
    list_ranges = {}

    def list_range(key, values):
        if key in list_ranges:
            return list_ranges[key]
        col = len(list_ranges) + 1
        lists.cell(1, col, key)
        for i, v in enumerate(values, start=2):
            lists.cell(i, col, v)
        letter = get_column_letter(col)
        ref = f"{quote_sheetname('Lists')}!${letter}$2:${letter}${len(values) + 1}"
        list_ranges[key] = ref
        return ref

    status_letters = {}
    for sh in SHEETS:
        ws = wb.create_sheet(sh.name)
        headers = ["Action"] + [f.header + (" *" if f.req else "") for f in sh.fields] + STATUS_COLS
        hints = ["Add / Update / Add/Update / Delete / Skip"] + [_hint(f) for f in sh.fields] + \
                ["written by the tool", "", ""]
        for c, (h, hint) in enumerate(zip(headers, hints), start=1):
            cell = ws.cell(HEADER_ROW, c, h)
            if c == 1:
                _hdr(cell, NAVY)
            elif c > len(sh.fields) + 1:
                _hdr(cell, GREY)
            else:
                f = sh.fields[c - 2]
                _hdr(cell, RED if f.req else (TEAL if f.key else NAVY))
            hc = ws.cell(HINT_ROW, c, hint)
            hc.font = Font(italic=True, color="666666", size=9)
            hc.alignment = Alignment(wrap_text=True, vertical="top")
            hc.fill = PatternFill("solid", fgColor="F2F4F8")
        ws.row_dimensions[HEADER_ROW].height = 32
        ws.row_dimensions[HINT_ROW].height = 42
        ws.column_dimensions["A"].width = 12
        for i, f in enumerate(sh.fields, start=2):
            ws.column_dimensions[get_column_letter(i)].width = f.width
        n = len(sh.fields)
        st_letter = get_column_letter(n + 2)
        status_letters[sh.name] = st_letter
        ws.column_dimensions[st_letter].width = 11
        ws.column_dimensions[get_column_letter(n + 3)].width = 60
        ws.column_dimensions[get_column_letter(n + 4)].width = 17
        ws.freeze_panes = "B3"

        # dropdowns
        last = DATA_ROW + MAX_ROWS
        dv = DataValidation(type="list", formula1=list_range("Action", ACTIONS), allow_blank=True)
        dv.add(f"A{DATA_ROW}:A{last}")
        ws.add_data_validation(dv)
        for i, f in enumerate(sh.fields, start=2):
            values = None
            if f.kind == "bool":
                values, key = BOOL_CHOICES, "Yes/No"
            elif f.choices == "enum":
                enum_path = f.path
                if sh.kind == "relations" and f.path != "name":
                    enum_path = f"relatedRegions.relatedRegion.{f.path}"
                values = schema.enum_for(sh.op("add"), enum_path) if enum_path else None
                key = f"{sh.obj}.{f.path}"
            elif isinstance(f.choices, list):
                values, key = f.choices, f"{sh.obj}.{f.header}"
            if values:
                letter = get_column_letter(i)
                dv = DataValidation(type="list", formula1=list_range(key, values),
                                    allow_blank=True, showErrorMessage=False)
                dv.add(f"{letter}{DATA_ROW}:{letter}{last}")
                ws.add_data_validation(dv)

        # status colours
        rng = f"{st_letter}{DATA_ROW}:{st_letter}{last}"
        for text, color in (("Success", "C6EFCE"), ("Updated", "C6EFCE"), ("Deleted", "C6EFCE"),
                            ("Failed", "FFC7CE"), ("Exists", "FFEB9C"), ("Not found", "FFEB9C"),
                            ("Validated", "DDEBF7"), ("Skipped", "EDEDED")):
            ws.conditional_formatting.add(
                rng, CellIsRule(operator="equal", formula=[f'"{text}"'],
                                fill=PatternFill("solid", fgColor=color)))
        ws.conditional_formatting.add(
            f"A{DATA_ROW}:A{last}",
            FormulaRule(formula=[f'$A{DATA_ROW}="Delete"'], font=Font(color="C00000", bold=True)))

        for r, row in enumerate((sample or {}).get(sh.name, []), start=DATA_ROW):
            ws.cell(r, 1, row.get("Action", "Add"))
            for i, f in enumerate(sh.fields, start=2):
                if f.header in row:
                    ws.cell(r, i, row[f.header])

    _build_dashboard(dash, status_letters)
    lists.sheet_state = "hidden"
    wb.save(path)
    return path


def _hint(f):
    parts = []
    if f.hint:
        parts.append(f.hint)
    if f.default is not None:
        parts.append(f"default: {f.default}")
    if f.kind == "bool":
        parts.append("Yes/No")
    return " | ".join(parts)


def _build_dashboard(ws, status_letters):
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 34
    for col, w in zip("CDEFGHI", (30, 10, 10, 10, 10, 10, 60)):
        ws.column_dimensions[col].width = w
    t = ws.cell(1, 2, "CUCM Greenfield Deployment Workbook")
    t.font = Font(bold=True, size=18, color=NAVY)
    ws.cell(2, 2, "Fill the sheets left to right, then run the CUCM Greenfield Deployer "
                  "(deployer_app.py). Status columns are written back by the tool.").font = \
        Font(italic=True, color="555555")

    ws.cell(4, 2, "Cluster").font = Font(bold=True, size=12, color=TEAL)
    for i, (label, value) in enumerate((("CUCM Publisher", ""), ("AXL User", ""),
                                        ("AXL Version", "Auto")), start=5):
        ws.cell(i, 2, label).font = Font(bold=True)
        c = ws.cell(i, 3, value)
        c.border = BORDER
        c.fill = PatternFill("solid", fgColor="FFF9E6")
    ws.cell(8, 2, "Password is asked by the tool - never stored in this file.").font = \
        Font(italic=True, size=9, color="888888")

    top = 10
    heads = ["Sheet (deploy order)", "Configures", "Rows", "Success", "Failed", "Exists",
             "Pending", "Notes"]
    for c, h in enumerate(heads, start=2):
        _hdr(ws.cell(top, c, h), NAVY)
    r = top + 1
    seen = set()
    for stage, name in DEPLOY_PLAN:
        if stage != "main" or name in seen:
            continue
        seen.add(name)
        sh = SHEET_BY_NAME[name]
        q = quote_sheetname(name)
        st = status_letters[name]
        link = ws.cell(r, 2, name)
        link.hyperlink = f"#{q}!A1"
        link.font = Font(color="0563C1", underline="single")
        ws.cell(r, 3, sh.obj)
        ws.cell(r, 4, f"=COUNTA({q}!B{DATA_ROW}:B{DATA_ROW + MAX_ROWS})")
        rng = f"{q}!{st}{DATA_ROW}:{st}{DATA_ROW + MAX_ROWS}"
        ws.cell(r, 5, f'=COUNTIF({rng},"Success")+COUNTIF({rng},"Updated")'
                      f'+COUNTIF({rng},"Deleted")')
        ws.cell(r, 6, f'=COUNTIF({rng},"Failed")')
        ws.cell(r, 7, f'=COUNTIF({rng},"Exists")')
        ws.cell(r, 8, f"=D{r}-E{r}-F{r}-G{r}")
        ws.cell(r, 9, sh.about)
        for c in range(2, 10):
            ws.cell(r, c).border = BORDER
            if r % 2:
                ws.cell(r, c).fill = PatternFill("solid", fgColor="F2F4F8")
        r += 1
    _hdr(ws.cell(r, 2, "Total"), TEAL)
    for c in range(3, 10):
        _hdr(ws.cell(r, c), TEAL)
    for c in "DEFGH":
        ws[f"{c}{r}"] = f"=SUM({c}{top + 1}:{c}{r - 1})"
    ws.conditional_formatting.add(f"F{top + 1}:F{r}", CellIsRule(
        operator="greaterThan", formula=["0"], fill=PatternFill("solid", fgColor="FFC7CE")))

    r += 2
    ws.cell(r, 2, "How to use").font = Font(bold=True, size=12, color=TEAL)
    tips = [
        "Red headers are required for Add. Teal headers identify the object for Update/Delete.",
        "Action: Add (default when blank), Update, Add/Update (add, or update if it exists), "
        "Delete, Skip.",
        "List cells use ';' between items, e.g. Partitions 'PT_Internal;PT_Local'. Order matters.",
        "Directory numbers are written DN/Partition, e.g. '1001/PT_Internal'.",
        "Type <none> in a cell to clear that field on Update.",
        "Any AXL field can be added as an extra column named 'axl:<path>', "
        "e.g. 'axl:callForwardAll.destination'.",
        "Deletes run first, in reverse order. Re-running skips rows already marked Success.",
        "Use 'Rollback workbook' in the tool to create a file that deletes what was added.",
    ]
    for tip in tips:
        r += 1
        ws.cell(r, 2, "• " + tip)


def default_output_path(path):
    base, ext = os.path.splitext(path)
    stamp = dt.datetime.now().strftime("%Y%m%d_%H%M")
    return f"{base}_result_{stamp}{ext}"
